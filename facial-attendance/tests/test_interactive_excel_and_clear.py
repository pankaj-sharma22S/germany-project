"""Tests for interactive Excel endpoints, daily in-place updates, and data clearing."""

import json
from pathlib import Path
import numpy as np
import openpyxl
import pytest
from fastapi.testclient import TestClient

from app.api import app, services
from app.config import Config
from app.excel_service import AttendanceEvent, ExcelService


@pytest.fixture
def client():
    services.excel_service.clear_all_data()
    return TestClient(app)


def test_excel_sheet_data_endpoints(client):
    # Test attendance sheet endpoint
    res_att = client.get("/api/excel/data?sheet=attendance")
    assert res_att.status_code == 200
    data_att = res_att.json()
    assert "columns" in data_att
    assert "rows" in data_att
    assert "summary" in data_att
    assert data_att["file_name"] == "attendance.xlsx"

    # Test students sheet endpoint
    res_stu = client.get("/api/excel/data?sheet=students")
    assert res_stu.status_code == 200
    data_stu = res_stu.json()
    assert "columns" in data_stu
    assert "rows" in data_stu
    assert "summary" in data_stu
    assert data_stu["file_name"] == "students.xlsx"

    # Test invalid sheet
    res_inv = client.get("/api/excel/data?sheet=invalid")
    assert res_inv.status_code == 400


def test_excel_download_endpoints(client):
    res_att = client.get("/api/excel/download/attendance.xlsx")
    assert res_att.status_code == 200
    assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in res_att.headers["content-type"]

    res_stu = client.get("/api/excel/download/students.xlsx")
    assert res_stu.status_code == 200

    res_inv = client.get("/api/excel/download/other.xlsx")
    assert res_inv.status_code == 404


def test_daily_in_place_update(tmp_path: Path):
    reg_file = tmp_path / "students.xlsx"
    att_file = tmp_path / "attendance.xlsx"

    service = ExcelService(
        project_root=tmp_path,
        registry_file=reg_file,
        attendance_file=att_file,
        attendance_mode="daily",
    )

    # First event on 2026-10-14
    ev1 = AttendanceEvent(
        session_id="SES_01",
        student_id="101",
        name="Alice",
        date_str="2026-10-14",
        time_str="10:00:00",
        status="Present",
        match_score=0.65,
    )
    written, errors = service.write_attendance_batch([ev1])
    assert written == 1
    assert len(errors) == 0

    wb = openpyxl.load_workbook(att_file)
    ws = wb["Attendance"]
    assert ws.max_row == 2  # Header + 1 record
    assert ws.cell(row=2, column=5).value == "10:00:00"
    assert ws.cell(row=2, column=7).value == 0.65
    wb.close()

    # Second event on SAME day in a new session SES_02 -> should UPDATE IN-PLACE
    ev2 = AttendanceEvent(
        session_id="SES_02",
        student_id="101",
        name="Alice",
        date_str="2026-10-14",
        time_str="14:30:00",
        status="Present",
        match_score=0.78,
    )
    written2, errors2 = service.write_attendance_batch([ev2])
    assert written2 == 1
    assert len(errors2) == 0

    wb = openpyxl.load_workbook(att_file)
    ws = wb["Attendance"]
    assert ws.max_row == 2  # STILL ONLY 1 ROW! No duplicate row created!
    assert ws.cell(row=2, column=1).value == "SES_02"
    assert ws.cell(row=2, column=5).value == "14:30:00"
    assert ws.cell(row=2, column=7).value == 0.78
    wb.close()


def test_clear_all_data_endpoint(client):
    res = client.post("/api/admin/clear-data")
    assert res.status_code == 200
    data = res.json()
    assert data["success"] is True

    # Verify sheet data after clearing
    res_att = client.get("/api/excel/data?sheet=attendance")
    assert res_att.json()["summary"]["total_rows"] == 0

    res_stu = client.get("/api/excel/data?sheet=students")
    assert res_stu.json()["summary"]["total_rows"] == 0
