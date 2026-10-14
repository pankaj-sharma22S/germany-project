"""Unit tests for ExcelService."""

from pathlib import Path
import cv2
import numpy as np
import openpyxl
import pandas as pd
import pytest

from app.excel_service import AttendanceEvent, ExcelService, StudentRecord


@pytest.fixture
def temp_project(tmp_path: Path):
    """Create a temporary project folder structure with photos."""
    data_dir = tmp_path / "data"
    data_dir.mkdir()
    photos_dir = data_dir / "photos"
    photos_dir.mkdir()

    # Create dummy valid image files
    img1 = np.ones((100, 100, 3), dtype=np.uint8) * 120
    cv2.imwrite(str(photos_dir / "S101.jpg"), img1)
    cv2.imwrite(str(photos_dir / "S102.jpg"), img1)

    # Create a corrupt non-image file
    corrupt_file = photos_dir / "corrupted.jpg"
    corrupt_file.write_text("NOT_AN_IMAGE_FILE")

    return tmp_path


def test_read_valid_registry(temp_project: Path):
    reg_file = temp_project / "data" / "students.xlsx"
    att_file = temp_project / "data" / "attendance.xlsx"

    # Create valid students.xlsx
    df = pd.DataFrame(
        [
            {"student_id": "S101", "name": "Alice Smith", "photo_path": "photos/S101.jpg"},
            {"student_id": "S102", "name": "Bob Jones", "photo_path": "photos/S102.jpg"},
        ]
    )
    df.to_excel(reg_file, index=False)

    service = ExcelService(
        project_root=temp_project,
        registry_file=reg_file,
        attendance_file=att_file,
    )

    records, errors = service.read_registry()
    assert len(errors) == 0
    assert len(records) == 2
    assert records[0].student_id == "S101"
    assert records[0].name == "Alice Smith"
    assert records[0].photo_path.exists()
    assert records[1].student_id == "S102"


def test_registry_duplicate_ids(temp_project: Path):
    reg_file = temp_project / "data" / "students.xlsx"
    att_file = temp_project / "data" / "attendance.xlsx"

    df = pd.DataFrame(
        [
            {"student_id": "S101", "name": "Alice Smith", "photo_path": "photos/S101.jpg"},
            {"student_id": "S101", "name": "Alice Duplicate", "photo_path": "photos/S101.jpg"},
        ]
    )
    df.to_excel(reg_file, index=False)

    service = ExcelService(temp_project, reg_file, att_file)
    records, errors = service.read_registry()
    assert len(records) == 1
    assert any("Duplicate 'student_id' 'S101'" in err for err in errors)


def test_registry_missing_columns(temp_project: Path):
    reg_file = temp_project / "data" / "students.xlsx"
    att_file = temp_project / "data" / "attendance.xlsx"

    df = pd.DataFrame([{"student_id": "S101", "name": "Alice Smith"}])  # missing photo_path
    df.to_excel(reg_file, index=False)

    service = ExcelService(temp_project, reg_file, att_file)
    records, errors = service.read_registry()
    assert len(records) == 0
    assert any("Missing required columns" in err for err in errors)


def test_registry_missing_photo_file(temp_project: Path):
    reg_file = temp_project / "data" / "students.xlsx"
    att_file = temp_project / "data" / "attendance.xlsx"

    df = pd.DataFrame(
        [
            {"student_id": "S103", "name": "Charlie", "photo_path": "photos/non_existent.jpg"}
        ]
    )
    df.to_excel(reg_file, index=False)

    service = ExcelService(temp_project, reg_file, att_file)
    records, errors = service.read_registry()
    assert len(records) == 0
    assert any("Reference photo file does not exist" in err for err in errors)


def test_registry_corrupt_photo_file(temp_project: Path):
    reg_file = temp_project / "data" / "students.xlsx"
    att_file = temp_project / "data" / "attendance.xlsx"

    df = pd.DataFrame(
        [{"student_id": "S104", "name": "David", "photo_path": "photos/corrupted.jpg"}]
    )
    df.to_excel(reg_file, index=False)

    service = ExcelService(temp_project, reg_file, att_file)
    records, errors = service.read_registry()
    assert len(records) == 0
    assert any("is corrupted or in an unsupported format" in err for err in errors)


def test_attendance_write_and_duplicate_prevention(temp_project: Path):
    reg_file = temp_project / "data" / "students.xlsx"
    att_file = temp_project / "data" / "attendance.xlsx"

    service = ExcelService(temp_project, reg_file, att_file)

    events_batch1 = [
        AttendanceEvent(
            session_id="SES_01",
            student_id="S101",
            name="Alice",
            date_str="2026-10-14",
            time_str="14:00:00",
            status="Present",
            match_score=0.75,
        ),
        AttendanceEvent(
            session_id="SES_01",
            student_id="S102",
            name="Bob",
            date_str="2026-10-14",
            time_str="14:01:00",
            status="Present",
            match_score=0.68,
        ),
    ]

    written, errors = service.write_attendance_batch(events_batch1)
    assert written == 2
    assert len(errors) == 0
    assert att_file.exists()

    # Verify Excel contents with openpyxl
    wb = openpyxl.load_workbook(att_file)
    ws = wb["Attendance"]
    assert ws.max_row == 3  # Header + 2 rows
    assert ws.cell(row=2, column=2).value == "S101"
    assert ws.cell(row=3, column=2).value == "S102"
    wb.close()

    # Attempt to write S101 again in the SAME session -> should be prevented
    events_batch2 = [
        AttendanceEvent(
            session_id="SES_01",
            student_id="S101",
            name="Alice",
            date_str="2026-10-14",
            time_str="14:05:00",
            status="Present",
            match_score=0.80,
        ),
        # New student S103 in SES_01
        AttendanceEvent(
            session_id="SES_01",
            student_id="S103",
            name="Charlie",
            date_str="2026-10-14",
            time_str="14:06:00",
            status="Present",
            match_score=0.72,
        ),
        # S101 in a DIFFERENT session SES_02 -> should be allowed
        AttendanceEvent(
            session_id="SES_02",
            student_id="S101",
            name="Alice",
            date_str="2026-10-14",
            time_str="15:00:00",
            status="Present",
            match_score=0.79,
        ),
    ]

    written2, errors2 = service.write_attendance_batch(events_batch2)
    assert written2 == 2  # S103 in SES_01, and S101 in SES_02; S101 in SES_01 skipped
    assert len(errors2) == 0

    wb2 = openpyxl.load_workbook(att_file)
    ws2 = wb2["Attendance"]
    assert ws2.max_row == 5  # Header + 4 records
    wb2.close()
